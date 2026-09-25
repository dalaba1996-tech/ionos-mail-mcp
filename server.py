from starlette.middleware.base import BaseHTTPMiddleware
from starlette.responses import JSONResponse
import uvicorn
import os
import imaplib
import email
import hmac
from email.header import decode_header
from mcp.server.fastmcp import FastMCP
from mcp.server.transport_security import TransportSecuritySettings

mcp = FastMCP(
    "IONOS Mail",
    stateless_http=True,
    json_response=True,
    transport_security=TransportSecuritySettings(
        enable_dns_rebinding_protection=True,
        allowed_hosts=[
            "ionos-mail-mcp-1.onrender.com",
            "ionos-mail-mcp-1.onrender.com:*",
        ],
        allowed_origins=[
            "https://ionos-mail-mcp-1.onrender.com",
        ],
    ),
)

IMAP_HOST = os.environ.get("IMAP_HOST", "imap.ionos.fr")
IMAP_PORT = int(os.environ.get("IMAP_PORT", "993"))
EMAIL_USER = os.environ.get("EMAIL_USER")
EMAIL_PASSWORD = os.environ.get("EMAIL_PASSWORD")
MCP_API_KEY = os.environ.get("MCP_API_KEY")


def connect_imap():
    if not EMAIL_USER or not EMAIL_PASSWORD:
        raise RuntimeError("Identifiants IONOS non configurés.")

    mail = imaplib.IMAP4_SSL(IMAP_HOST, IMAP_PORT)
    mail.login(EMAIL_USER, EMAIL_PASSWORD)
    return mail


def decode_text(value):
    if not value:
        return ""

    parts = decode_header(value)
    result = ""

    for text, charset in parts:
        if isinstance(text, bytes):
            result += text.decode(charset or "utf-8", errors="replace")
        else:
            result += text

    return result


@mcp.tool()
def recent_emails(limit: int = 10) -> str:
    """Retourne les derniers emails reçus dans la boîte IONOS."""

    limit = max(1, min(limit, 50))

    mail = connect_imap()

    try:
        mail.select("INBOX", readonly=True)
        status, data = mail.search(None, "ALL")

        if status != "OK":
            return "Impossible de rechercher les emails."

        ids = data[0].split()[-limit:]
        results = []

        for message_id in reversed(ids):
            status, msg_data = mail.fetch(
                message_id,
                "(BODY.PEEK[HEADER.FIELDS (FROM SUBJECT DATE)])"
            )

            if status != "OK":
                continue

            raw = next(
                (
                    item[1]
                    for item in msg_data
                    if isinstance(item, tuple)
                ),
                None,
            )

            if not raw:
                continue

            msg = email.message_from_bytes(raw)

            results.append(
    f"ID: {message_id.decode()}\n"
    f"De: {decode_text(msg.get('From'))}\n"
    f"Objet: {decode_text(msg.get('Subject'))}\n"
    f"Date: {msg.get('Date', '')}"
)

        return "\n\n---\n\n".join(results) or "Aucun email trouvé."

    finally:
        try:
            mail.logout()
        except Exception:
            pass


@mcp.tool()
def search_emails(query: str, limit: int = 10) -> str:
    """Recherche des emails contenant un texte dans l'objet."""

    limit = max(1, min(limit, 50))
    mail = connect_imap()

    try:
        mail.select("INBOX", readonly=True)

        safe_query = query.replace('"', "")
        status, data = mail.search(
            None,
            "SUBJECT",
            f'"{safe_query}"'
        )

        if status != "OK":
            return "Recherche impossible."

        ids = data[0].split()[-limit:]
        results = []

        for message_id in reversed(ids):
            status, msg_data = mail.fetch(
                message_id,
                "(BODY.PEEK[HEADER.FIELDS (FROM SUBJECT DATE)])"
            )

            if status != "OK":
                continue

            raw = next(
                (
                    item[1]
                    for item in msg_data
                    if isinstance(item, tuple)
                ),
                None,
            )

            if not raw:
                continue

            msg = email.message_from_bytes(raw)

            results.append(
    f"ID: {message_id.decode()}\n"
    f"De: {decode_text(msg.get('From'))}\n"
    f"Objet: {decode_text(msg.get('Subject'))}\n"
    f"Date: {msg.get('Date', '')}"
)
        return "\n\n---\n\n".join(results) or "Aucun email trouvé."

    finally:
        try:
            mail.logout()
        except Exception:
            pass

@mcp.tool()
def read_email(message_id: str) -> str:
    """Lit le contenu complet d'un email IONOS à partir de son identifiant."""

    mail = connect_imap()

    try:
        mail.select("INBOX", readonly=True)

        status, msg_data = mail.fetch(message_id, "(RFC822)")

        if status != "OK":
            return "Impossible de récupérer cet email."

        raw_email = next(
            (
                item[1]
                for item in msg_data
                if isinstance(item, tuple)
            ),
            None,
        )

        if not raw_email:
            return "Email introuvable."

        msg = email.message_from_bytes(raw_email)

        sender = decode_text(msg.get("From"))
        recipient = decode_text(msg.get("To"))
        subject = decode_text(msg.get("Subject"))
        date = msg.get("Date", "")

        body = ""

        if msg.is_multipart():
            for part in msg.walk():
                content_type = part.get_content_type()
                disposition = str(
                    part.get("Content-Disposition", "")
                ).lower()

                # On ignore les pièces jointes pour l'instant
                if "attachment" in disposition:
                    continue

                if content_type == "text/plain":
                    payload = part.get_payload(decode=True)

                    if payload:
                        charset = part.get_content_charset() or "utf-8"
                        body = payload.decode(
                            charset,
                            errors="replace"
                        )
                        break

        else:
            payload = msg.get_payload(decode=True)

            if payload:
                charset = msg.get_content_charset() or "utf-8"
                body = payload.decode(
                    charset,
                    errors="replace"
                )

        if not body.strip():
            body = (
                "Le message ne contient pas de version texte "
                "directement lisible."
            )

        return (
            f"De: {sender}\n"
            f"À: {recipient}\n"
            f"Objet: {subject}\n"
            f"Date: {date}\n\n"
            f"Contenu:\n{body.strip()}"
        )

    finally:
        try:
            mail.logout()
        except Exception:
            pass
            
@mcp.tool()
def list_attachments(message_id: str) -> str:
    """Liste les pièces jointes présentes dans un email IONOS."""

    mail = connect_imap()

    try:
        mail.select("INBOX", readonly=True)

        status, msg_data = mail.fetch(message_id, "(RFC822)")

        if status != "OK":
            return "Impossible de récupérer cet email."

        raw_email = next(
            (
                item[1]
                for item in msg_data
                if isinstance(item, tuple)
            ),
            None,
        )

        if not raw_email:
            return "Email introuvable."

        msg = email.message_from_bytes(raw_email)

        attachments = []

        for index, part in enumerate(msg.walk()):
            filename = part.get_filename()

            if filename:
                filename = decode_text(filename)

                attachments.append(
                    f"ID pièce jointe: {index}\n"
                    f"Nom: {filename}\n"
                    f"Type: {part.get_content_type()}\n"
                    f"Taille: "
                    f"{len(part.get_payload(decode=True) or b'')} octets"
                )

        if not attachments:
            return "Cet email ne contient aucune pièce jointe."

        return "\n\n---\n\n".join(attachments)

    finally:
        try:
            mail.logout()
        except Exception:
            pass
@mcp.tool()
def read_attachment(message_id: str, attachment_id: int) -> str:
    """Lit le contenu d'une pièce jointe PDF, DOCX, XLSX, CSV ou TXT."""

    import io
    import csv
    from pypdf import PdfReader
    from docx import Document
    from openpyxl import load_workbook

    mail = connect_imap()

    try:
        mail.select("INBOX", readonly=True)
        status, msg_data = mail.fetch(message_id, "(RFC822)")

        if status != "OK":
            return "Impossible de récupérer cet email."

        raw_email = next(
            (item[1] for item in msg_data if isinstance(item, tuple)),
            None,
        )

        if not raw_email:
            return "Email introuvable."

        msg = email.message_from_bytes(raw_email)

        parts = list(msg.walk())

        if attachment_id < 0 or attachment_id >= len(parts):
            return "Identifiant de pièce jointe invalide."

        part = parts[attachment_id]
        filename = part.get_filename()

        if not filename:
            return "Cette partie du message n'est pas une pièce jointe."

        filename = decode_text(filename)
        data = part.get_payload(decode=True)

        if not data:
            return "La pièce jointe est vide."

        # Limite de sécurité : 10 Mo
        if len(data) > 10 * 1024 * 1024:
            return "Pièce jointe trop volumineuse (limite : 10 Mo)."

        extension = filename.lower().rsplit(".", 1)[-1]

        if extension == "pdf":
            reader = PdfReader(io.BytesIO(data))
            text = "\n\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        elif extension == "docx":
            document = Document(io.BytesIO(data))
            text = "\n".join(
                paragraph.text
                for paragraph in document.paragraphs
            )

        elif extension == "xlsx":
            workbook = load_workbook(
                io.BytesIO(data),
                read_only=True,
                data_only=True
            )

            output = []

            for sheet in workbook.worksheets:
                output.append(f"Feuille: {sheet.title}")

                for row in sheet.iter_rows(values_only=True):
                    output.append(
                        " | ".join(
                            "" if value is None else str(value)
                            for value in row
                        )
                    )

            text = "\n".join(output)

        elif extension == "csv":
            decoded = data.decode("utf-8-sig", errors="replace")
            rows = csv.reader(io.StringIO(decoded))

            text = "\n".join(
                " | ".join(row)
                for row in rows
            )

        elif extension in ("txt", "md"):
            text = data.decode("utf-8", errors="replace")

        else:
            return (
                f"Format non pris en charge : {filename}. "
                "Formats acceptés : PDF, DOCX, XLSX, CSV, TXT et MD."
            )

        # Évite d'envoyer une quantité énorme de texte à Claude
        max_chars = 100_000

        if len(text) > max_chars:
            text = (
                text[:max_chars]
                + "\n\n[Contenu tronqué après 100 000 caractères]"
            )

        return (
            f"Fichier: {filename}\n\n"
            f"{text.strip() or '[Aucun texte extractible]'}"
        )

    except Exception as exc:
        return f"Impossible de lire la pièce jointe : {type(exc).__name__}"

    finally:
        try:
            mail.logout()
        except Exception:
            pass
            
class APIKeyMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        if request.url.path.startswith("/mcp"):
            provided_key = request.headers.get("X-API-Key", "")

            if not MCP_API_KEY or not hmac.compare_digest(
                provided_key,
                MCP_API_KEY
            ):
                return JSONResponse(
                    {"error": "Unauthorized"},
                    status_code=401
                )

        return await call_next(request)


app = mcp.streamable_http_app()
app.add_middleware(APIKeyMiddleware)


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "10000"))

    uvicorn.run(
        app,
        host="0.0.0.0",
        port=port
    )
